import json
import os
import time

import openpyxl
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError, expect

from test.playwright.e2e.utils.helpers import (
    open_new_profile_form,
    close_profile_form_if_open,
    dismiss_onboarding_modals,
    dismiss_tour_if_present,
    poll_sample_accession,
    run_django_shell,
    upload_manifest,
)

from common.utils.helpers import get_env

TESTFILES_DIR = os.path.normpath(os.path.join(os.path.dirname(__file__), '..', '..', 'testfiles'))
SAMPLE_MANIFEST = os.path.join(TESTFILES_DIR, 'COPO_FAANG_manifest_success1.xlsx')
# Holds only what a customer fills in. Everything EI/Sapio generates per
# project (Project ID, budget holder, JIRA ticket, Sample ID, EI sample name,
# plate and well) changes every run, so the test downloads the manifest COPO
# pre-fills for this run's profile and copies these answers into it.
EDP_LIMS_MANIFEST = os.path.join(TESTFILES_DIR, 'edp_lims_test_manifest_tissue_success1.xlsx')

FAANG_SAMPLE_CHECKLIST = 'COPO_FAANG'
# The EI_EDP schema's checklists_filter maps sample type 'Tissue' to this
# checklist, and the LIMS page offers only that one for a Tissue profile.
EDP_CHECKLIST = 'version_tissue'

# The manifest's study-level answers, keyed by their EI_EDP schema label.
STUDY_CELLS = {'Email Address': 'E6', 'Sample Return': 'L8', 'Health and Safety': 'L9'}
# The sample sheet's EI-only block (Sample ID ... Project ID) ends where this
# column starts; every column from here rightwards is the customer's.
FIRST_SUBMITTER_LABEL = 'Submitter Sample Reference'
# These three are VLOOKUP array formulas into the sample_metadata sheet,
# which COPO fills from the profile's own ENA samples. COPO reads the upload
# with data_only=True (EDPSchemasSpreadsheet.loadManifest), i.e. it reads
# Excel's cached formula results -- and a workbook saved by openpyxl has none,
# so the formulas would read back empty. Resolve them here instead.
SAMPLE_METADATA_LOOKUPS = {
    'NCBI Taxon ID': 'taxon_id',
    'Scientific Name (Latin)': 'scientific_name',
    'BioSample Accession': 'biosampleAccession',
}

# EDP profile form values agreed for the Sapio test instance. The JIRA ticket
# number is deliberately not here: it becomes the Sapio ProjectName, so each
# run makes its own (see edp_related_enq()) to find its project in
# Sapio unambiguously.
EDP_PROFILE_FIELDS = {
    'sample_type': 'Tissue',
    'library_type': 'Illumina premade library',
    'container_type': 'Plate',
    'no_of_samples': 1,
    'budget_holder': 'copo-playwright-test-user',
    
    # The project team's development/test email address (which is set in
    # docker-compose.playwright.yaml) will receive emails about the 
    # shared COPO profile which is sent during the tests.
    # COPO skips sending emails when an email is not set up (like during CI). The
    # profile creation will still proceed.
    # The reserved `.invalid` fallback is never emailed.
    'customer_emails': get_env('PROJECT_TEAM_EMAIL') or 'copo-playwright@example.invalid',
}

def edp_related_enq():
    return f'copo-playwright-test-{int(time.time())}'

def header_columns(sheet, row):
    return {str(cell.value).strip(): cell.column for cell in sheet[row] if cell.value not in (None, '')}


def find_header_row(sheet):
    return next(row[0].row for row in sheet.iter_rows() if row[0].value == 'Sample ID')


def manifest_sample_ids(path):
    # The Sample IDs listed on a downloaded manifest's sample sheet
    sheet = openpyxl.load_workbook(path)['sample']
    header_row = find_header_row(sheet)
    column = header_columns(sheet, header_row)['Sample ID']
    return [
        sheet.cell(row, column).value
        for row in range(header_row + 1, sheet.max_row + 1)
        if sheet.cell(row, column).value not in (None, '')
    ]


def build_edp_manifest(downloaded_path, patched_path, accession, submitter_reference=None):
    # Fills the one sample row of a downloaded manifest with the answers file.
    # Returns (sapio_sample_id, expected) where expected maps each schema
    # label to the value this manifest submits for it.
    #
    # submitter_reference replaces the answers file's Submitter Sample
    # Reference, so that one ENA sample can be submitted as several Sapio
    # samples, each under its own customer name.
    downloaded = openpyxl.load_workbook(downloaded_path)
    answers = openpyxl.load_workbook(EDP_LIMS_MANIFEST)

    sample_sheet, answer_sheet = downloaded['sample'], answers['sample']
    expected = {}
    for label, cell in STUDY_CELLS.items():
        sample_sheet[cell] = expected[label] = answer_sheet[cell].value

    header_row = find_header_row(sample_sheet)
    answer_header_row = find_header_row(answer_sheet)
    columns = header_columns(sample_sheet, header_row)
    answer_columns = header_columns(answer_sheet, answer_header_row)
    data_row, answer_row = header_row + 1, answer_header_row + 1

    sapio_sample_id = sample_sheet.cell(data_row, columns['Sample ID']).value
    assert sapio_sample_id, 'Downloaded manifest has no pre-filled Sapio Sample ID'
    assert sample_sheet.cell(data_row + 1, columns['Sample ID']).value in (None, ''), (
        'Downloaded manifest has more than one sample row; the answers file describes one'
    )

    reference = answer_sheet.cell(answer_row, answer_columns[FIRST_SUBMITTER_LABEL]).value
    metadata_sheet = downloaded['sample_metadata']
    metadata_header = [cell.value for cell in metadata_sheet[1]]
    metadata = {
        values[0]: dict(zip(metadata_header, values))
        for values in metadata_sheet.iter_rows(min_row=2, values_only=True)
        if values[0]
    }
    assert reference in metadata, (
        f'Sample {reference!r} is missing from the downloaded sample_metadata sheet '
        f"(found {sorted(metadata)}); the ENA sample's name must match the manifest's "
        f'Submitter Sample Reference'
    )
    assert metadata[reference]['biosampleAccession'] == accession

    first_submitter_column = columns[FIRST_SUBMITTER_LABEL]
    for label, column in columns.items():
        if column < first_submitter_column:
            continue
        if label in SAMPLE_METADATA_LOOKUPS:
            value = metadata[reference][SAMPLE_METADATA_LOOKUPS[label]]
        elif label == FIRST_SUBMITTER_LABEL and submitter_reference:
            value = submitter_reference
        else:
            value = answer_sheet.cell(answer_row, answer_columns[label]).value
        sample_sheet.cell(data_row, column).value = expected[label] = value

    # The index sheet's Sample ID is the Earlham Institute's (EI's) (Sapio's) ID for the sample, which
    # only exists once the profile has been saved.
    index_sheet, answer_index_sheet = downloaded['index'], answers['index']
    index_columns = header_columns(index_sheet, 1)
    for label, answer_column in header_columns(answer_index_sheet, 1).items():
        if label in ('Project ID', 'Sample ID'):
            continue
        index_sheet.cell(2, index_columns[label]).value = answer_index_sheet.cell(2, answer_column).value
    index_sheet.cell(2, index_columns['Sample ID']).value = sapio_sample_id

    downloaded.save(patched_path)
    return sapio_sample_id, expected


def submit_sample_to_ena(page, profile_id):
    # Uploads the FAANG sample manifest, submits its sample to ENA's dev
    # sandbox and returns the real BioSample accession. The LIMS manifest
    # looks each sample's taxon and accession up from these ENA samples.
    page.goto(f'/copo/copo_sample/view/general/{profile_id}/general_sample')
    page.wait_for_load_state('networkidle')
    dismiss_tour_if_present(page)

    page.select_option('#checklist_id', FAANG_SAMPLE_CHECKLIST)
    page.click('.new-general-sample-spreadsheet-template')
    page.wait_for_load_state('networkidle')
    page.wait_for_timeout(1000)
    upload_manifest(page, SAMPLE_MANIFEST)
    try:
        expect(page.locator('#save_sample_button')).to_be_enabled(timeout=60000)
    except AssertionError as error:
        # parse_sample_spreadsheet answers a failed validation with an empty
        # 400; the reasons only reach the page via the websocket.
        raise AssertionError(f"Sample manifest rejected: {page.locator('#sample_info').inner_text()}") from error
    page.click('#save_sample_button')

    page.wait_for_timeout(1000)
    dismiss_tour_if_present(page)
    page.locator('#sample_table tbody tr').first.click()
    page.click(".copo-dt[data-action='submit_sample']")
    # A fresh test user has no ENA credentials of its own, so COPO offers
    # its default; with valid credentials the submission just proceeds.
    try:
        page.locator('#repo_cred_modal .cred-modal-default').wait_for(state='visible', timeout=5000)
    except PlaywrightTimeoutError:
        pass
    else:
        page.click('#repo_cred_modal .cred-modal-default')

    return poll_sample_accession(profile_id)


def download_edp_manifest(page, profile_id, path):
    # The blank LIMS manifest, pre-filled with the project's samples still
    # to be submitted
    response = page.request.get(
        f'/copo/copo_single_cell/download_initial_manifest/{profile_id}/EI_EDP/{EDP_CHECKLIST}'
    )
    assert response.ok, f'Manifest download failed: {response.status}'
    path.write_bytes(response.body())
    return path


def upload_edp_manifest(page, profile_id, path):
    # Opens the LIMS page afresh and uploads the manifest. Returns the
    # validation response; the reasons for a rejection only reach the page
    # (#singlecell_info) via the websocket.
    page.goto(f'/copo/copo_single_cell/view/{profile_id}/EI_EDP/lims_integration')
    page.wait_for_load_state('networkidle')
    dismiss_tour_if_present(page)

    page.select_option('#checklist_id', EDP_CHECKLIST)
    page.click('.new-singlecell-spreadsheet-template')
    page.wait_for_load_state('networkidle')
    page.wait_for_timeout(1000)
    with page.expect_response(lambda r: 'parse_singlecell_spreadsheet' in r.url, timeout=120000) as response_info:
        upload_manifest(page, str(path))
    return response_info.value


def save_edp_manifest(page):
    expect(page.locator('#save_singlecell_button')).to_be_enabled(timeout=60000)
    page.click('#save_singlecell_button')
    page.wait_for_load_state('networkidle')


def submit_edp_to_sapio(page):
    # Same table mechanics as the single-cell journey: the study tab's table
    # is `singlecell_table_study`, the toolbar is cloned into every tab (so
    # scope the click to the study table's own button), and the study tab
    # auto-selects its first row after each refresh -- clicking the row
    # ourselves would race that.
    #
    # The Sapio write happens inside this request (no Celery), so its own
    # feedback is the outcome. Returns that feedback.
    page.wait_for_timeout(1000)
    dismiss_tour_if_present(page)
    page.wait_for_function(
        "() => $.fn.dataTable.isDataTable('#singlecell_table_study') && "
        "$('#singlecell_table_study').DataTable().rows({selected: true}).count() > 0",
        timeout=30000,
    )
    with page.expect_response(lambda r: 'copo_forms' in r.url, timeout=120000) as response_info:
        page.click(".copo-dt[data-action='submit_singlecell_sapio'][data-table='singlecell_table_study']")
    response = response_info.value
    assert response.status == 200, f'submit_singlecell_sapio request failed: {response.status} {response.text()[:500]}'
    return response.json().get('action_feedback', {})


def poll_sapio_study_status(profile_id, study_id, timeout=60, interval=5):
    # submit_edp_to_sapio() records the outcome on the singlecell study as
    # status_sapio / accession_sapio / error_sapio (update_component_status
    # suffixes each key with the repository).
    code = (
        'from src.apps.copo_single_cell_submission.utils.da import Singlecell\n'
        f"records = Singlecell(profile_id={profile_id!r}).execute_query({{'profile_id': {profile_id!r}}})\n"
        'study = {}\n'
        'for r in records:\n'
        "    for s in r.get('components', {}).get('study', []):\n"
        f"        if s.get('study_id') == {study_id!r}:\n"
        '            study = s\n'
        "print('POLL_RESULT:' + '|'.join(["
        "str(study.get('status_sapio', '')), str(study.get('accession_sapio', '')), str(study.get('error_sapio', ''))"
        ']))\n'
    )
    deadline = time.time() + timeout
    status = ''
    while time.time() < deadline:
        marker_line = next(
            (line for line in run_django_shell(code).splitlines() if line.startswith('POLL_RESULT:')),
            'POLL_RESULT:||',
        )
        status, accession, error = marker_line[len('POLL_RESULT:'):].split('|', 2)
        if status == 'accepted' and accession:
            return accession
        if status == 'rejected' or error:
            raise AssertionError(f'Sapio submission recorded as failed: status={status!r} error={error!r}')
        time.sleep(interval)
    raise AssertionError(f'Sapio submission status never recorded (last status={status!r})')


def copo_sample_ids(profile_id):
    # The Sample IDs on the profile's saved LIMS record in COPO
    code = (
        'from src.apps.copo_single_cell_submission.utils.da import Singlecell\n'
        f"records = Singlecell(profile_id={profile_id!r}).execute_query({{'profile_id': {profile_id!r}}})\n"
        "ids = [str(s.get('sample_id', '')) for r in records for s in r.get('components', {}).get('sample', [])]\n"
        "print('SAMPLE_IDS:' + ','.join(ids))\n"
    )
    marker_line = next(line for line in run_django_shell(code).splitlines() if line.startswith('SAMPLE_IDS:'))
    return [sample_id for sample_id in marker_line[len('SAMPLE_IDS:'):].split(',') if sample_id]

def create_edp_profile(
    page,
    title,
    description,
    *,
    sample_type,
    library_type,
    container_type,
    no_of_samples,
    related_enq,
    budget_holder,
    customer_emails,
):
    # Saving an EDP profile is not local-only: post_save_edp_profile()
    # creates a Sapio project (plus its samples and plates) and stores the
    # project's identifier back on the profile, and the 'Sample type'
    # dropdown's options are themselves loaded from Sapio when the form
    # opens. Every caller therefore needs the Sapio test instance reachable.
    #
    # Returns (profile_card, profile_id, sapio_project_id).
    open_new_profile_form(page, 'ei_edp')

    page.fill('#copo\\.profile\\.title', title)
    page.fill('#copo\\.profile\\.description', description)
    page.select_option('#copo\\.profile\\.sample_type', sample_type)
    page.select_option('#copo\\.profile\\.library_type', library_type)
    # <option> values are lower-case ('plate'/'tube'); select by the visible
    # label so callers can pass what the user sees.
    page.select_option('#copo\\.profile\\.container_type', label=container_type)
    page.fill('#copo\\.profile\\.no_of_samples', str(no_of_samples))
    page.fill('#copo\\.profile\\.related_enq', related_enq)
    page.fill('#copo\\.profile\\.budget_holder', budget_holder)
    page.fill('#copo\\.profile\\.customer_emails', customer_emails)

    # Check the save's response rather than only the card appearing: a
    # Sapio sync failure still saves the COPO profile (and renders its card)
    # but reports a non-success status here.
    with page.expect_response(
        lambda r: 'copo_forms' in r.url and r.request.method == 'POST', timeout=120000
    ) as response_info:
        page.click('#btnFormSave')
    feedback = response_info.value.json().get('action_feedback', {})
    assert feedback.get('status') == 'success', f'EDP profile save failed: {feedback}'

    profile_card = page.locator('.panel-heading', has_text=title)
    # The Sapio sync happens inside the save request, and the card's
    # re-render has been seen to take well past create_biodata_profile()'s
    # 15s budget (one exploratory run never showed it within 60s at all,
    # the next showed it within 8s).
    expect(profile_card).to_be_visible(timeout=60000)

    close_profile_form_if_open(page)

    # The card's title <span> carries the profile id; the Sapio project id
    # sits in the card's hidden 'View more details' popover content, so
    # read it with text_content() rather than waiting for it to be visible.
    profile_id = profile_card.locator('.row-title span').first.get_attribute('id')
    sapio_project_id = (page.locator(f'#sapio_project_id_{profile_id}').text_content() or '').strip()
    return profile_card, profile_id, sapio_project_id


def edit_edp_profile_sample_count(page, title, no_of_samples):
    # Changes the number of samples on an existing EDP profile and saves.
    # Returns the save's feedback. On a rejected save the form stays open
    # with the message shown in it, so it is left open for the caller to
    # check; on success it is closed.
    page.goto('/copo')
    dismiss_tour_if_present(page)
    dismiss_onboarding_modals(page)
    dismiss_tour_if_present(page)

    profile_card = page.locator('.panel-heading', has_text=title)
    profile_card.locator('.profile-ellipsis').click()
    page.click('#editProfileBtn')

    sample_count = page.locator('#copo\\.profile\\.no_of_samples')
    # The form is built server-side, and its Sample type options come from Sapio
    expect(sample_count).to_be_visible(timeout=60000)
    sample_count.fill(str(no_of_samples))

    with page.expect_response(
        lambda r: 'copo_forms' in r.url and r.request.method == 'POST', timeout=120000
    ) as response_info:
        page.click('#btnFormSave')
    feedback = response_info.value.json().get('action_feedback', {})
    if feedback.get('status') != 'error':
        close_profile_form_if_open(page)
    return feedback

def read_sapio_project(project_id, project_fields=(), sample_fields=()):
    # Reads directly from Sapio and not from COPO's copy in Mongo database, 
    # so that a test can confirm what actually was recorded/saved in the LIMS. 
    # Uses the adapter's sample loader so 'samples under this 
    # project' means exactly what COPO means by it when submitting.
    #
    # Returns {'exists': bool, 'project': {field: value}, 'samples': [{field: value}, ...]}.
    code = (
        'import json\n'
        'from src.apps.ei_edp.utils.lims.sapio.adapter import SapioAdapter\n'
        'from src.apps.ei_edp.utils.lims.sapio.datamanager import Sapio\n'
        f'project_id = {project_id!r}\n'
        'records = Sapio().dataRecordManager.query_data_records(\n'
        "    data_type_name='Project', data_field_name='C_ProjectIdentifier', value_list=[project_id]\n"
        ').result_list\n'
        "result = {'exists': bool(records), 'project': {}, 'samples': []}\n"
        'if records:\n'
        f"    result['project'] = {{f: records[0].get_field_value(f) for f in {list(project_fields)!r}}}\n"
        "    result['samples'] = [\n"
        f'        {{f: s.get_field_value(f) for f in {list(sample_fields)!r}}}\n'
        '        for s in SapioAdapter()._load_project_samples(project_id)\n'
        '    ]\n'
        "print('SAPIO_RESULT:' + json.dumps(result, default=str))\n"
    )
    marker_line = next(
        line for line in run_django_shell(code).splitlines() if line.startswith('SAPIO_RESULT:')
    )
    return json.loads(marker_line[len('SAPIO_RESULT:'):])
